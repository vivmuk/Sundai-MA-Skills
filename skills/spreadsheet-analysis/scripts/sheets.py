#!/usr/bin/env python3
"""Inspect, summarise and write spreadsheets — with the data-quality checks first.

The mess in a clinical spreadsheet is not cosmetic; it changes the answer.
Inconsistent category spellings split a count until the signal disappears.
`inspect` is the most valuable subcommand and should be run first, always.

    S=skills/spreadsheet-analysis/scripts/sheets.py
    python3 $S inspect   --file insights.csv
    python3 $S summarise --file insights.csv --by setting,country
    python3 $S dedupe    --file insights.csv --key contact_id --show-conflicts
    python3 $S write     --spec table.json --out tracker.xlsx

openpyxl is optional — without it, `write` emits CSV plus a markdown table.
"""
from __future__ import annotations
import argparse, csv, json, re, sys
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path

DATE_RE = re.compile(r"^\s*\d{1,4}[-/ ]\d{1,2}[-/ ]\d{1,4}\s*$")
NUM_RE = re.compile(r"^-?\d+(\.\d+)?%?$")


def _have(m: str) -> bool:
    try:
        __import__(m); return True
    except ImportError:
        return False


def load(path: Path, sheet: str | None) -> tuple[list[str], list[dict], list[str]]:
    notes = []
    if path.suffix.lower() in (".csv", ".tsv"):
        lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
        skipped = 0
        while lines and lines[0].lstrip().startswith("#"):
            skipped += 1; lines.pop(0)
        if skipped:
            notes.append(f"Skipped {skipped} leading comment line(s).")
        delim = "\t" if path.suffix.lower() == ".tsv" else ","
        rows = list(csv.DictReader(lines, delimiter=delim))
        return (list(rows[0].keys()) if rows else []), rows, notes
    if not _have("openpyxl"):
        raise SystemExit(
            "Reading .xlsx needs openpyxl:  pip install openpyxl\n"
            "Alternative with no install: re-save the sheet as CSV and pass that."
        )
    import openpyxl
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    ws = wb[sheet] if sheet and not sheet.isdigit() else wb[wb.sheetnames[int(sheet)-1 if sheet else 0]]
    raw = [[("" if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
    raw = [r for r in raw if any(x.strip() for x in r)]
    if len(raw) > 1 and sum(1 for c in raw[1] if c.strip()) > sum(1 for c in raw[0] if c.strip()):
        notes.append("Row 2 has more populated cells than row 1 — this may be a "
                     "TWO-ROW HEADER. Check before trusting the column names.")
    hdr = raw[0]
    return hdr, [dict(zip(hdr, r)) for r in raw[1:]], notes


def similar_groups(values: list[str], threshold: float = 0.82) -> list[list[str]]:
    """Candidate category variants that would split a count if left unmerged."""
    uniq = sorted({v.strip() for v in values if v.strip()})
    groups, used = [], set()
    for i, a in enumerate(uniq):
        if a in used:
            continue
        g = [a]
        for b in uniq[i+1:]:
            if b in used:
                continue
            na, nb = a.lower().replace(" ", ""), b.lower().replace(" ", "")
            if na == nb or SequenceMatcher(None, na, nb).ratio() >= threshold:
                g.append(b); used.add(b)
        if len(g) > 1:
            groups.append(g); used.add(a)
    return groups


def cmd_inspect(cols, rows, notes) -> int:
    print(f"{len(rows)} data rows × {len(cols)} columns\n")
    for n in notes:
        print(f"  NOTE  {n}")
    problems = 0
    for c in cols:
        vals = [str(r.get(c, "")) for r in rows]
        nonblank = [v for v in vals if v.strip()]
        blank = len(vals) - len(nonblank)
        kinds = set()
        for v in nonblank[:400]:
            kinds.add("number" if NUM_RE.match(v) else
                      "date-as-text" if DATE_RE.match(v) else "text")
        uniq = len(set(nonblank))
        print(f"\n  {c}")
        print(f"    {len(nonblank)} populated, {blank} blank, {uniq} distinct, "
              f"types: {', '.join(sorted(kinds)) or '—'}")
        if blank and blank < len(vals):
            print(f"    ! {blank} blank cell(s) — merged cells read as one value "
                  f"plus blanks")
            problems += 1
        if "date-as-text" in kinds and "text" in kinds:
            print("    ! dates stored as text mixed with other text — sorting will "
                  "be wrong, and day-first vs month-first is silent")
            problems += 1
        if 1 < uniq <= 60:
            groups = similar_groups(nonblank)
            for g in groups:
                counts = Counter(nonblank)
                total = sum(counts[x] for x in g)
                print(f"    ! likely one concept split across {len(g)} spellings "
                      f"(total n={total}): {', '.join(repr(x) for x in g)}")
                problems += 1
    if problems:
        print(f"\n{problems} data-quality issue(s). Split categories are the one "
              f"that changes the answer — the frequency that mattered disappears\n"
              f"into several small ones. Use medical-terminology-mapping, and "
              f"publish the collapsing decisions alongside any count.")
    else:
        print("\nNo structural problems detected. That is not a guarantee — check "
              "the first five rows by eye before trusting the header.")
    return 0


def cmd_summarise(cols, rows, by: str) -> int:
    keys = [k.strip() for k in by.split(",") if k.strip()]
    for k in keys:
        if k not in cols:
            print(f"No such column: {k}. Available: {', '.join(cols)}", file=sys.stderr)
            return 2
    counts = Counter(tuple(str(r.get(k, "")).strip() for k in keys) for r in rows)
    print(f"n = {len(rows)} records\n")
    print("| " + " | ".join(keys) + " | n | % |")
    print("| " + " | ".join("---" for _ in keys) + " | ---: | ---: |")
    for combo, n in counts.most_common():
        pct = 100 * n / len(rows) if rows else 0
        print("| " + " | ".join(combo) + f" | {n} | {pct:.1f} |")
    print(f"\nDenominator is {len(rows)} records. If several rows describe the same "
          f"contact, this is a RECORD-level count, not a person-level one — say "
          f"which in the deliverable.")
    return 0


def cmd_dedupe(cols, rows, key: str, show_conflicts: bool) -> int:
    if key not in cols:
        print(f"No such column: {key}. Available: {', '.join(cols)}", file=sys.stderr)
        return 2
    groups = defaultdict(list)
    for r in rows:
        groups[str(r.get(key, "")).strip()].append(r)
    repeats = {k: v for k, v in groups.items() if len(v) > 1}
    print(f"{len(rows)} rows · {len(groups)} distinct {key} · "
          f"{len(repeats)} appearing more than once\n")
    conflicts = 0
    if show_conflicts:
        for k, group in sorted(repeats.items()):
            varying = [c for c in cols if c != key
                       and len({str(r.get(c, "")).strip() for r in group}) > 1]
            stable_like = [c for c in varying
                           if c.lower() in ("setting", "country", "region", "site",
                                            "institution", "specialty")]
            if stable_like:
                conflicts += 1
                print(f"  ! {key}={k} has conflicting {', '.join(stable_like)}")
                for c in stable_like:
                    print(f"      {c}: {sorted({str(r.get(c,'')).strip() for r in group})}")
    if conflicts:
        print(f"\n{conflicts} identifier(s) carry conflicting attributes that should "
              f"be properties of the entity, not of the record.\nEither the key is "
              f"not stable, or those fields are entered per-interaction. Until it is "
              f"resolved, entity-level frequency claims cannot be trusted — report "
              f"at record level and say so.")
    else:
        print("No attribute conflicts detected on the key.")
    print(f"\nRepeated {key} values are not necessarily duplicates: several people "
          f"reporting the same contact is an ECHO; different contacts saying the "
          f"same thing is a PATTERN. That distinction is usually the whole analysis.")
    return 0


def cmd_write(spec: dict, out: Path) -> int:
    cols, rows = spec["columns"], spec["rows"]
    if _have("openpyxl"):
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook(); ws = wb.active
        ws.title = spec.get("sheet_name", "Data")[:31]
        ws.append([spec.get("draft_marking",
                            "DRAFT — NOT FOR EXTERNAL USE. REQUIRES QUALIFIED MEDICAL REVIEW.")])
        ws["A1"].font = Font(bold=True, color="993300")
        ws.append([])
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(row=3, column=c)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EEF2F6")
        for r in rows:
            ws.append(list(r))
        ws.freeze_panes = "A4"
        for i, col in enumerate(cols, 1):
            width = max([len(str(col))] + [len(str(r[i-1])) for r in rows if i <= len(r)])
            ws.column_dimensions[chr(64 + i) if i <= 26 else "AA"].width = min(width + 3, 60)
        out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(out))
        print(f"Wrote {out} ({len(rows)} rows). Header frozen at row 3.")
        return 0

    csv_path, md_path = out.with_suffix(".csv"), out.with_suffix(".md")
    out.parent.mkdir(parents=True, exist_ok=True)
    with csv_path.open("w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh); w.writerow(cols); w.writerows(rows)
    md = ["**DRAFT — NOT FOR EXTERNAL USE. REQUIRES QUALIFIED MEDICAL REVIEW.**", "",
          "| " + " | ".join(cols) + " |", "| " + " | ".join("---" for _ in cols) + " |"]
    md += ["| " + " | ".join(str(c) for c in r) + " |" for r in rows]
    md_path.write_text("\n".join(md) + "\n", encoding="utf-8")
    print(f"""⚠ DEGRADED OUTPUT — openpyxl is not available in this environment.

   Delivered:      {csv_path}  (the data)
                   {md_path}  (readable table)
   Not delivered:  {out}
   To get it:      pip install openpyxl
                   python3 {Path(__file__).name} write --spec <spec> --out {out}

   The data is complete. Formulas, formatting and multiple sheets were lost;
   nothing else was.
""")
    return 0


# A worked spec, used by --example and by scripts/selftest_fallbacks.py. The
# denominator column is deliberate: a tracker of percentages with no n is the
# commonest way a field-insight summary becomes uninterpretable.
EXAMPLE_SPEC = {
    "sheet_name": "Insight tracker Q2",
    "draft_marking": "DRAFT — NOT FOR EXTERNAL USE. REQUIRES QUALIFIED MEDICAL REVIEW.",
    "columns": ["Insight ID", "Theme", "Contacts raising", "n screened",
                "% of contacts", "Confidence", "Owner", "Decision it informs"],
    "rows": [
        ["INS-001", "Infection prophylaxis practice varies by centre",
         14, 40, "35.0%", "Moderate", "Medical Director, Haematology",
         "Whether to develop prophylaxis guidance for the field"],
        ["INS-002", "Step-up dosing burden limits community uptake",
         9, 40, "22.5%", "Moderate", "Head of Field Medical",
         "Community-setting evidence generation priority"],
        ["INS-003", "Uncertainty about sequencing after BCMA exposure",
         6, 40, "15.0%", "Low — weak signal, 6 contacts",
         "Evidence Generation Lead", "Whether to fund the sequencing analysis"],
    ],
}


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("inspect", "summarise", "dedupe"):
        p = sub.add_parser(name)
        p.add_argument("--file", required=True)
        p.add_argument("--sheet")
        if name == "summarise":
            p.add_argument("--by", required=True)
        if name == "dedupe":
            p.add_argument("--key", required=True)
            p.add_argument("--show-conflicts", action="store_true")
    w = sub.add_parser("write")
    w.add_argument("--spec")
    w.add_argument("--out", default="table.xlsx")
    w.add_argument("--example", action="store_true",
                   help="print a worked spec to stdout and exit")
    args = ap.parse_args()

    if args.cmd == "write":
        if args.example:
            print(json.dumps(EXAMPLE_SPEC, indent=2))
            return 0
        if not args.spec:
            print("write needs --spec (or --example to see one)", file=sys.stderr)
            return 2
        return cmd_write(json.loads(Path(args.spec).read_text(encoding="utf-8")),
                         Path(args.out))
    path = Path(args.file)
    if not path.exists():
        print(f"No such file: {path}", file=sys.stderr); return 2
    cols, rows, notes = load(path, args.sheet)
    if args.cmd == "inspect":
        return cmd_inspect(cols, rows, notes)
    if args.cmd == "summarise":
        return cmd_summarise(cols, rows, args.by)
    return cmd_dedupe(cols, rows, args.key, args.show_conflicts)


if __name__ == "__main__":
    raise SystemExit(main())
