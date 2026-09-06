#!/usr/bin/env python3
"""Formal Medical Affairs correspondence, with the compliance gates enforced.

A letter is a durable, forwardable, quotable record, and the warm helpful
register is exactly the one in which an unsubstantiated claim slips through.

    S=skills/medical-correspondence/scripts/letter.py
    python3 $S --example dhcp > letter.json
    python3 $S --spec letter.json --out letter.docx

python-docx is optional; falls back to markdown.
"""
from __future__ import annotations
import argparse, json, sys
from pathlib import Path

DRAFT = "DRAFT — NOT FOR EXTERNAL USE. REQUIRES QUALIFIED MEDICAL REVIEW."
TYPES = ["dhcp", "investigator-response", "agency-brief",
         "advisory-invitation", "author-cover"]

EXAMPLES = {
    "dhcp": {
        "type": "dhcp",
        "subject": "Important safety information: NORVANTIB (norvantimab) — "
                   "updated guidance on infection prophylaxis",
        "recipient": "Dear Healthcare Professional",
        "sender": {"name": "Dr [Name]", "role": "Medical Director",
                   "org": "Nordvant Biopharma (fictional)"},
        "date": "2026-08-15",
        "purpose": ("This letter informs you of updated guidance on infection "
                    "prophylaxis for patients receiving NORVANTIB, following review "
                    "of post-marketing reports of serious infections."),
        "body": [
            {"heading": "What has changed",
             "text": "The product information has been updated to recommend "
                     "antimicrobial prophylaxis for all patients during and for "
                     "3 months after treatment."},
            {"heading": "Basis for this change",
             "text": "Cumulative review identified serious infections in patients "
                     "not receiving prophylaxis at a higher rate than observed in "
                     "the pivotal trial, in which prophylaxis was mandated."},
        ],
        "actions": [
            "Initiate antimicrobial prophylaxis in all patients starting NORVANTIB.",
            "Monitor immunoglobulin levels every 4 weeks during treatment.",
            "Review patients currently on treatment without prophylaxis.",
        ],
        "ae_reporting": ("Report suspected adverse reactions to [national reporting "
                         "scheme] or to Nordvant Biopharma at [contact]. Reporting "
                         "helps identify new safety information."),
        "contact": "Medical Information: [phone] · [email]",
        "approval_status": ("NORVANTIB is approved for adults with relapsed or "
                            "refractory multiple myeloma after ≥4 prior lines of therapy."),
    },
    "advisory-invitation": {
        "type": "advisory-invitation",
        "subject": "Invitation: advisory board on sequencing after BCMA-directed therapy",
        "recipient": "Dear Dr [Name]",
        "sender": {"name": "Dr [Name]", "role": "Therapeutic Area Medical Lead",
                   "org": "Nordvant Biopharma (fictional)"},
        "date": "2026-08-15",
        "purpose": ("We are seeking advice on a question we cannot answer from "
                    "our own data: how clinicians should sequence therapy after "
                    "BCMA-directed treatment, and what evidence would change "
                    "current practice."),
        "body": [
            {"heading": "Why you", "text": "Your prospective work on discontinuation "
             "and your published position on duration of therapy address exactly "
             "the question we are unable to answer."},
            {"heading": "The commitment", "text": "One 4-hour meeting on [date] in "
             "[location], plus approximately 2 hours of pre-reading."},
            {"heading": "Compensation", "text": "Fair market value honorarium of "
             "[amount], plus reasonable travel expenses."},
        ],
        "transparency": ("Payments will be disclosed under applicable transparency "
                         "legislation, including the US Sunshine Act and the EFPIA "
                         "Disclosure Code."),
        "contact": "[name] · [email]",
    },
}

REQUIRED = {
    "dhcp": [("purpose", "state what changed and what the reader must do, first"),
             ("actions", "specific actions for the recipient, not general advice"),
             ("ae_reporting", "adverse event reporting instructions with the actual route"),
             ("contact", "a named contact for questions")],
    "advisory-invitation": [
        ("purpose", "the question you are asking them"),
        ("transparency", "the transparency-reporting consequence of accepting"),
    ],
}


def _have(m: str) -> bool:
    try:
        __import__(m); return True
    except ImportError:
        return False


def check(spec: dict) -> list[str]:
    out = []
    kind = spec.get("type")
    if kind not in TYPES:
        out.append(f"type must be one of {', '.join(TYPES)}")
        return out
    for field, why in REQUIRED.get(kind, []):
        if not spec.get(field):
            out.append(f"{kind}: missing '{field}' — {why}")
    text = json.dumps(spec).lower()
    if any(w in text for w in ("superior", "best-in-class", "well-tolerated", "proven")):
        out.append("promotional language detected — a letter is a durable, "
                   "forwardable record and this is where the boundary gets crossed")
    if kind in ("dhcp", "investigator-response") and not spec.get("approval_status"):
        out.append(f"{kind}: no approval_status statement")
    return out


def render_markdown(spec: dict) -> str:
    s = spec.get("sender", {})
    L = [f"**{DRAFT}**", "", f"{s.get('org','')}", f"{spec.get('date','')}", "",
         f"**{spec.get('subject','')}**", "", spec.get("recipient", "Dear Colleague"), "",
         spec.get("purpose", ""), ""]
    for b in spec.get("body", []):
        L += [f"### {b['heading']}", "", b["text"], ""]
    if spec.get("actions"):
        L += ["### What we are asking you to do", ""]
        L += [f"{i}. {a}" for i, a in enumerate(spec["actions"], 1)] + [""]
    for key, head in [("approval_status", "Approval status"),
                      ("transparency", "Transparency reporting"),
                      ("ae_reporting", "Reporting adverse reactions")]:
        if spec.get(key):
            L += [f"### {head}", "", spec[key], ""]
    if spec.get("contact"):
        L += ["### Contact", "", spec["contact"], ""]
    L += ["Yours sincerely", "", s.get("name", ""), s.get("role", ""), s.get("org", ""),
          "", f"**{DRAFT}**", ""]
    return "\n".join(L)


def build_docx(spec: dict, out: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    warn = doc.add_paragraph().add_run(DRAFT)
    warn.bold = True; warn.font.size = Pt(9); warn.font.color.rgb = RGBColor(0x99, 0x33, 0)
    s = spec.get("sender", {})
    doc.add_paragraph(s.get("org", "")); doc.add_paragraph(spec.get("date", ""))
    doc.add_paragraph()
    doc.add_paragraph().add_run(spec.get("subject", "")).bold = True
    doc.add_paragraph(spec.get("recipient", "Dear Colleague"))
    doc.add_paragraph(spec.get("purpose", ""))
    for b in spec.get("body", []):
        doc.add_heading(b["heading"], level=2); doc.add_paragraph(b["text"])
    if spec.get("actions"):
        doc.add_heading("What we are asking you to do", level=2)
        for a in spec["actions"]:
            doc.add_paragraph(a, style="List Number")
    for key, head in [("approval_status", "Approval status"),
                      ("transparency", "Transparency reporting"),
                      ("ae_reporting", "Reporting adverse reactions")]:
        if spec.get(key):
            doc.add_heading(head, level=2); doc.add_paragraph(spec[key])
    if spec.get("contact"):
        doc.add_heading("Contact", level=2); doc.add_paragraph(spec["contact"])
    doc.add_paragraph(); doc.add_paragraph("Yours sincerely")
    for line in (s.get("name", ""), s.get("role", ""), s.get("org", "")):
        doc.add_paragraph(line)
    r = doc.add_paragraph().add_run(DRAFT)
    r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x99, 0x33, 0)
    out.parent.mkdir(parents=True, exist_ok=True); doc.save(str(out))


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--spec"); ap.add_argument("--out", default="letter.docx")
    ap.add_argument("--example", choices=sorted(EXAMPLES))
    args = ap.parse_args()
    if args.example:
        print(json.dumps(EXAMPLES[args.example], indent=2)); return 0
    if not args.spec:
        ap.print_help(); return 2
    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
    problems = check(spec)
    for p in problems:
        print(f"  BLOCKED  {p}")
    if problems:
        return 1
    out = Path(args.out)
    if _have("docx"):
        build_docx(spec, out); print(f"Wrote {out}")
    else:
        md = out.with_suffix(".md"); md.parent.mkdir(parents=True, exist_ok=True)
        md.write_text(render_markdown(spec), encoding="utf-8")
        print(f"""⚠ DEGRADED OUTPUT — python-docx is not available.

   Delivered:      {md}  (complete letter, all compliance sections)
   Not delivered:  {out}
   To get it:      pip install python-docx && python3 {Path(__file__).name} --spec {args.spec} --out {out}

   The content is complete. Only Word formatting was lost.
""")
    if spec.get("type") == "dhcp":
        print("\nDHCP letters have a review pathway that usually includes regulatory "
              "and often the regulator itself. Never send one on the strength of "
              "this script alone.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
