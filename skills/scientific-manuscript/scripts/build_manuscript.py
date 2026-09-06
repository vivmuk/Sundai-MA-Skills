#!/usr/bin/env python3
"""Render a structured scientific manuscript to .docx, and gap-check it.

Structure follows the reporting guideline, so the guideline is a required field
rather than an afterthought — choosing it after drafting means restructuring.

    S=skills/scientific-manuscript/scripts/build_manuscript.py

    python3 $S --example > ms.json
    python3 $S --spec ms.json --out manuscript.docx
    python3 $S --spec ms.json --checklist consort   # gap-check before submission

Requires python-docx:  pip install python-docx
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

DRAFT = "DRAFT — NOT FOR SUBMISSION. REQUIRES QUALIFIED MEDICAL REVIEW."

GUIDELINES = {
    "consort": "CONSORT 2010 — randomised trials",
    "spirit": "SPIRIT 2013 — trial protocols",
    "prisma": "PRISMA 2020 — systematic reviews and meta-analyses",
    "strobe": "STROBE — observational studies",
    "record": "RECORD — routinely-collected health data",
    "care": "CARE — case reports",
    "cheers": "CHEERS 2022 — economic evaluations",
    "stard": "STARD 2015 — diagnostic accuracy",
    "tripod": "TRIPOD — prediction models",
}

# The items most often missing at submission, per guideline. Not the full
# checklist — the full checklist belongs to EQUATOR. These are the ones that
# cause avoidable revision rounds.
HIGH_RISK_ITEMS = {
    "consort": [
        ("design named in title and abstract", "title", "abstract"),
        ("trial registration number and registry", "registration"),
        ("protocol availability", "protocol"),
        ("sample size determination", "methods"),
        ("randomisation: sequence generation, concealment, implementation", "methods"),
        ("blinding, and of whom", "methods"),
        ("statistical methods including multiplicity handling", "methods"),
        ("participant flow diagram", "results"),
        ("primary outcome with absolute AND relative effect sizes and CIs", "results"),
        ("all harms, not only those above a threshold", "results"),
        ("limitations, addressing sources of potential bias", "discussion"),
        ("funding and the funder's role", "funding"),
    ],
    "prisma": [
        ("identified as a systematic review in the title", "title"),
        ("protocol registration (PROSPERO)", "registration"),
        ("full search strategy for every database, with dates", "methods"),
        ("selection process and number of reviewers", "methods"),
        ("risk of bias tool named", "methods"),
        ("certainty assessment (GRADE)", "methods"),
        ("flow diagram with reconciling numbers", "results"),
        ("excluded studies WITH REASONS", "results"),
        ("certainty of evidence per outcome", "results"),
    ],
    "strobe": [
        ("design named in title and abstract", "title", "abstract"),
        ("eligibility criteria and sources of participants", "methods"),
        ("all variables, including confounders and effect modifiers", "methods"),
        ("how confounding was addressed", "methods"),
        ("handling of missing data", "methods"),
        ("participant numbers at each stage", "results"),
        ("unadjusted AND confounder-adjusted estimates with CIs", "results"),
        ("limitations, including direction and magnitude of potential bias", "discussion"),
    ],
    "care": [
        ("'case report' in the title", "title"),
        ("de-identified patient information", "methods"),
        ("timeline of the episode", "results"),
        ("patient perspective", "discussion"),
        ("informed consent obtained", "ethics"),
    ],
}

EXAMPLE = {
    "title": (
        "A randomised, open-label, phase 3 trial of NORVANTIB versus standard "
        "of care in relapsed or refractory multiple myeloma"
    ),
    "short_title": "NORVANTIB in relapsed/refractory myeloma",
    "guideline": "consort",
    "registration": "ClinicalTrials.gov NCT00000000; registered 2021-03-14, before first enrolment",
    "protocol": "Protocol and statistical analysis plan available as supplementary material",
    "authors": [
        {"name": "A Example", "affiliation": "Fictional Cancer Centre, Springfield",
         "icmje": ["design", "interpretation", "drafting", "approval", "accountable"]},
        {"name": "B Sample", "affiliation": "Nordvant Biopharma (fictional)",
         "icmje": ["analysis", "revision", "approval", "accountable"]},
    ],
    "corresponding": "A Example, a.example@example.invalid",
    "writing_support": (
        "Medical writing support was provided by [name], [agency], and funded by "
        "Nordvant Biopharma."
    ),
    "funding": (
        "Funded by Nordvant Biopharma (fictional). The funder participated in "
        "study design, data analysis, and the decision to submit."
    ),
    "conflicts": "A Example reports advisory fees from [x]. B Sample is an employee of Nordvant Biopharma.",
    "data_sharing": "De-identified participant data available on request subject to a data access agreement.",
    "ethics": "Approved by the institutional review board at each site; all participants gave written informed consent.",
    "keywords": ["multiple myeloma", "bispecific antibody", "randomised trial"],
    "abstract": {
        "Background": "Treatment options after triple-class exposure are limited.",
        "Methods": "We randomly assigned 558 patients 1:1 to NORVANTIB or standard of care. The primary endpoint was progression-free survival.",
        "Results": "Median PFS was 20.4 months (95% CI 17.1-24.8) versus 10.2 months (8.9-12.4); HR 0.58 (95% CI 0.47-0.72).",
        "Conclusions": "NORVANTIB prolonged progression-free survival compared with standard of care in this population.",
        "Registration": "NCT00000000",
    },
    "sections": {
        "Introduction": "What is known. What is not known. What this study asked.",
        "Methods": "Design, setting, participants, interventions, outcomes, sample size, randomisation, blinding, statistical methods including the testing hierarchy.",
        "Results": "Participant flow, baseline characteristics, primary outcome, secondary outcomes, harms.",
        "Discussion": "Principal findings. Comparison with existing evidence. Limitations. Implications.",
        "Conclusion": "Proportionate to the design.",
    },
    "references": [
        "Example A, Sample B, et al. Prior work in this population. J Example Med. 2024;12(3):100-110. doi:10.0000/example",
    ],
}


def check(spec: dict) -> tuple[list[str], list[str]]:
    """Returns (blocking, advisory)."""
    blocking: list[str] = []
    advisory: list[str] = []

    if not spec.get("guideline"):
        blocking.append(
            "no reporting guideline declared. Structure follows the guideline — "
            f"choose one of: {', '.join(sorted(GUIDELINES))}"
        )
    elif spec["guideline"] not in GUIDELINES:
        blocking.append(
            f"unknown guideline '{spec['guideline']}'. "
            f"Options: {', '.join(sorted(GUIDELINES))}"
        )

    if not spec.get("conflicts"):
        blocking.append(
            "no conflicts of interest statement. Required by every journal and "
            "by ICMJE."
        )
    if not spec.get("funding"):
        blocking.append(
            "no funding statement. GPP 2022 requires the sponsor's role in "
            "design, analysis and the decision to publish to be stated."
        )

    authors = spec.get("authors") or []
    if not authors:
        blocking.append("no authors listed")
    for a in authors:
        crit = set(a.get("icmje") or [])
        # ICMJE requires all four. Criterion 1 is satisfied by design OR
        # acquisition/analysis/interpretation; criterion 2 by drafting OR revision.
        c1 = crit & {"design", "conception", "acquisition", "analysis", "interpretation"}
        c2 = crit & {"drafting", "revision"}
        c3 = "approval" in crit
        c4 = "accountable" in crit
        missing = []
        if not c1:
            missing.append("1 (conception/design or acquisition/analysis/interpretation)")
        if not c2:
            missing.append("2 (drafting or critical revision)")
        if not c3:
            missing.append("3 (final approval)")
        if not c4:
            missing.append("4 (accountability)")
        if missing:
            blocking.append(
                f"author '{a.get('name', '?')}' does not meet ICMJE criteria "
                f"{', '.join(missing)}. All four are required; contributors who do "
                f"not meet them belong in the acknowledgements with their "
                f"contribution described."
            )

    if spec.get("guideline") in ("consort", "spirit") and not spec.get("registration"):
        blocking.append(
            "trial registration is required for CONSORT/SPIRIT, and must predate "
            "first enrolment."
        )

    if not spec.get("writing_support"):
        advisory.append(
            "no writing_support statement. If professional medical writing "
            "support was used it MUST be disclosed with the writer, affiliation "
            "and funder — undisclosed support is ghostwriting. If none was used, "
            "state that explicitly."
        )
    if not spec.get("data_sharing"):
        advisory.append("no data sharing statement — required by most journals and by ICMJE")
    if not spec.get("ethics"):
        advisory.append("no ethics approval / consent statement")
    if not spec.get("references"):
        advisory.append("no references listed")

    return blocking, advisory


def checklist(guideline: str, spec: dict) -> list[tuple[str, bool]]:
    """Crude presence check of the high-risk items against the drafted text."""
    items = HIGH_RISK_ITEMS.get(guideline, [])
    haystack = json.dumps(spec).lower()
    out = []
    for item in items:
        label, *fields = item
        # Look for the section content being non-trivially present.
        found = False
        for f in fields:
            val = spec.get(f) or (spec.get("sections") or {}).get(f.title(), "")
            if isinstance(val, dict):
                val = " ".join(str(v) for v in val.values())
            if isinstance(val, (list, tuple)):
                val = " ".join(str(v) for v in val)
            if str(val).strip():
                found = True
        out.append((label, found))
    return out


def _have(module: str) -> bool:
    """Inline capability check, kept small so a copied directory still works."""
    try:
        __import__(module)
        return True
    except ImportError:
        return False


def render_markdown(spec: dict) -> str:
    """Tier 3: the full manuscript as markdown.

    Every element a journal requires survives — title page, authorship with
    affiliations, structured abstract, IMRaD body, the statements block, and
    references. Only Word styling is lost, and most journals accept a plain
    submission anyway.
    """
    out = [f"**{DRAFT}**", "", f"# {spec['title']}", ""]
    if spec.get("short_title"):
        out += [f"*Running head: {spec['short_title']}*", ""]

    out += ["## Authors", ""]
    for a in spec.get("authors", []):
        crit = ", ".join(a.get("icmje", []))
        out.append(f"- **{a['name']}** — {a.get('affiliation', '')}"
                   + (f"  <sub>ICMJE: {crit}</sub>" if crit else ""))
    if spec.get("corresponding"):
        out += ["", f"Corresponding author: {spec['corresponding']}"]
    out += ["", f"**Reporting guideline:** {GUIDELINES.get(spec.get('guideline', ''), '—')}"]
    if spec.get("registration"):
        out.append(f"**Registration:** {spec['registration']}")
    if spec.get("protocol"):
        out.append(f"**Protocol:** {spec['protocol']}")
    if spec.get("keywords"):
        out.append(f"**Keywords:** {', '.join(spec['keywords'])}")
    out += ["", "---", ""]

    if spec.get("abstract"):
        out += ["## Abstract", ""]
        for head, body in spec["abstract"].items():
            out.append(f"**{head}.** {body}")
            out.append("")
        out += ["> Every number in this abstract must also appear in the Results "
                "with its confidence interval.", "", "---", ""]

    for heading, body in (spec.get("sections") or {}).items():
        out += [f"## {heading}", "", str(body), ""]

    out += ["---", ""]
    for label, key in [
        ("Ethics approval and consent", "ethics"),
        ("Funding", "funding"),
        ("Role of the funder", "funder_role"),
        ("Medical writing support", "writing_support"),
        ("Conflicts of interest", "conflicts"),
        ("Data sharing", "data_sharing"),
        ("Acknowledgements", "acknowledgements"),
    ]:
        if spec.get(key):
            out += [f"### {label}", "", spec[key], ""]

    if spec.get("references"):
        out += ["## References", ""]
        out += [f"{i}. {r}" for i, r in enumerate(spec["references"], 1)]
        out.append("")

    out += ["", f"**{DRAFT}**", ""]
    return "\n".join(out)


def build(spec: dict, out: Path) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def para(text="", *, size=11, bold=False, italic=False, colour=None,
             align=None, space_after=8):
        p = doc.add_paragraph()
        if align:
            p.alignment = align
        run = p.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold
        run.italic = italic
        if colour:
            run.font.color.rgb = colour
        p.paragraph_format.space_after = Pt(space_after)
        return p

    WARN = RGBColor(0x99, 0x33, 0x00)

    # ---- title page -----------------------------------------------------
    para(DRAFT, size=10, bold=True, colour=WARN)
    para(spec["title"], size=16, bold=True, space_after=14)
    if spec.get("short_title"):
        para(f"Running head: {spec['short_title']}", size=10, italic=True)

    para("Authors", size=12, bold=True, space_after=4)
    for a in spec.get("authors", []):
        para(f"{a['name']} — {a.get('affiliation', '')}", size=10, space_after=2)
    if spec.get("corresponding"):
        para(f"Corresponding author: {spec['corresponding']}", size=10)

    if spec.get("keywords"):
        para("Keywords: " + ", ".join(spec["keywords"]), size=10, italic=True)

    para(f"Reporting guideline: {GUIDELINES.get(spec.get('guideline', ''), '—')}",
         size=10, italic=True)
    if spec.get("registration"):
        para(f"Registration: {spec['registration']}", size=10)
    if spec.get("protocol"):
        para(f"Protocol: {spec['protocol']}", size=10)

    doc.add_page_break()

    # ---- abstract -------------------------------------------------------
    if spec.get("abstract"):
        doc.add_heading("Abstract", level=1)
        for head, body in spec["abstract"].items():
            p = doc.add_paragraph()
            r = p.add_run(f"{head}. ")
            r.bold = True
            r.font.size = Pt(10)
            r2 = p.add_run(str(body))
            r2.font.size = Pt(10)
        para(
            "Every number in this abstract must also appear in the Results with "
            "its confidence interval. The abstract is the part most people read "
            "and the part most likely to overclaim.",
            size=8, italic=True, colour=WARN,
        )
        doc.add_page_break()

    # ---- body -----------------------------------------------------------
    for heading, body in (spec.get("sections") or {}).items():
        doc.add_heading(heading, level=1)
        for chunk in str(body).split("\n\n"):
            para(chunk)

    # ---- statements -----------------------------------------------------
    doc.add_page_break()
    for label, key in [
        ("Ethics approval and consent", "ethics"),
        ("Funding", "funding"),
        ("Role of the funder", "funder_role"),
        ("Medical writing support", "writing_support"),
        ("Conflicts of interest", "conflicts"),
        ("Data sharing", "data_sharing"),
        ("Acknowledgements", "acknowledgements"),
    ]:
        if spec.get(key):
            doc.add_heading(label, level=2)
            para(spec[key], size=10)

    if spec.get("references"):
        doc.add_heading("References", level=1)
        for i, ref in enumerate(spec["references"], 1):
            para(f"{i}. {ref}", size=10, space_after=4)

    para("")
    para(DRAFT, size=10, bold=True, colour=WARN)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--spec")
    ap.add_argument("--out", default="manuscript.docx")
    ap.add_argument("--example", action="store_true")
    ap.add_argument("--checklist", choices=sorted(HIGH_RISK_ITEMS),
                    help="gap-check the high-risk items for this guideline")
    args = ap.parse_args()

    if args.example:
        print(json.dumps(EXAMPLE, indent=2))
        return 0
    if not args.spec:
        ap.print_help()
        return 2

    path = Path(args.spec)
    if not path.exists():
        print(f"No such file: {path}", file=sys.stderr)
        return 2
    spec = json.loads(path.read_text(encoding="utf-8"))

    blocking, advisory = check(spec)
    for b in blocking:
        print(f"  BLOCKED   {b}")
    for a in advisory:
        print(f"  advisory  {a}")

    if args.checklist:
        print(f"\nHigh-risk items — {GUIDELINES[args.checklist]}\n")
        rows = checklist(args.checklist, spec)
        for label, found in rows:
            print(f"  [{'x' if found else ' '}] {label}")
        missing = sum(1 for _, f in rows if not f)
        print(
            f"\n{len(rows) - missing}/{len(rows)} present. This is a presence check "
            f"of the items that most often cause an avoidable revision round — it "
            f"is not the full checklist. Complete the official one from the "
            f"EQUATOR Network before submission."
        )
        return 1 if (blocking or missing) else 0

    if blocking:
        print(f"\n{len(blocking)} blocking problem(s). These are integrity "
              f"requirements, not formatting preferences.")
        return 1

    out = Path(args.out)
    NEXT = (
        "\nStill required before submission: citation-integrity over every "
        "reference, deliverable-quality-review for spin and overclaiming, and "
        "the full EQUATOR checklist for your design."
    )

    if _have("docx"):
        build(spec, out)
        print(f"Wrote {out}")
        print(NEXT)
        return 0

    md_path = out.with_suffix(".md")
    md_path.parent.mkdir(parents=True, exist_ok=True)
    md_path.write_text(render_markdown(spec), encoding="utf-8")
    print(f"""⚠ DEGRADED OUTPUT — python-docx is not available in this environment.

   Delivered:      {md_path}  (complete manuscript, IMRaD, statements, references)
   Not delivered:  {out}
   To get it:      pip install python-docx
                   python3 {Path(__file__).name} --spec {args.spec} --out {out}

   The content is complete. Title page, authorship with ICMJE criteria, the
   structured abstract, the full body, every statement and the reference list
   all survived. Only Word styling was lost — and most journals accept a plain
   submission.
""")
    print(NEXT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
