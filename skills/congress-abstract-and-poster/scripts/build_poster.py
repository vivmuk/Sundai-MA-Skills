#!/usr/bin/env python3
"""Build a scientific poster at the physical size the congress specifies.

A poster is read standing up, from about two metres, in a crowded hall, by
someone deciding in three seconds whether to stop. This builder enforces the
minimum legible font sizes for the dimensions you give it and refuses to render
body text below the threshold — the most common poster defect, and one that
cannot be fixed at the venue.

    S=skills/congress-abstract-and-poster/scripts/build_poster.py

    python3 $S --example > poster.json
    python3 $S --spec poster.json --out poster.pptx

There is no universal poster size. Confirm width, height, and orientation with
the congress AND the printer before building.

Requires python-pptx:  pip install python-pptx
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

# Below these, a reader at two metres cannot read the poster. They are not
# style preferences; they are the difference between a poster that works and
# one that is walked past.
MIN_BODY_PT = 24
MIN_HEADING_PT = 36
MIN_TITLE_PT = 72

EXAMPLE = {
    "width_in": 48,
    "height_in": 36,
    "title": "NORVANTIB in triple-class-exposed relapsed/refractory multiple myeloma",
    "authors": "A Example, B Sample, C Illustration",
    "affiliations": "Fictional Cancer Centre, Springfield; Nordvant Biopharma (fictional)",
    "abstract_number": "Abstract 1234",
    "takeaway": (
        "A 63.0% response rate was observed in a single-arm study. "
        "Randomised comparison is required."
    ),
    "columns": [
        {
            "sections": [
                {"heading": "Background",
                 "body": "Treatment options after triple-class exposure are limited. "
                         "This study evaluated NORVANTIB in this population."},
                {"heading": "Methods",
                 "body": "Single-arm, phase 1/2 study. Patients (N=165) with relapsed "
                         "or refractory multiple myeloma and ≥ 3 prior lines received "
                         "NORVANTIB weekly after step-up dosing. Primary endpoint: "
                         "overall response rate."},
            ]
        },
        {
            "sections": [
                {"heading": "Results",
                 "body": "Overall response rate 63.0% (95% CI 55.2–70.4).\n"
                         "Median duration of response 18.4 months (14.9–NE).\n"
                         "Median follow-up 14.1 months."},
                {"heading": "Safety",
                 "body": "Cytokine release syndrome 72.1% (grade ≥3: 0.6%).\n"
                         "Neutropenia 70.9% (grade ≥3: 64.2%).\n"
                         "Adverse events led to discontinuation in 4.2%."},
            ]
        },
        {
            "sections": [
                {"heading": "Conclusions",
                 "body": "A response rate of 63.0% was observed in this single-arm "
                         "study. No comparative inference is available; randomised "
                         "comparison against current standard of care is required."},
                {"heading": "Disclosures",
                 "body": "Funded by Nordvant Biopharma (fictional). A Example reports "
                         "advisory fees. B Sample is an employee of the sponsor."},
            ]
        },
    ],
    "approval_status": (
        "NORVANTIB is investigational. Efficacy and safety have not been established."
    ),
    "footer": "SYNTHETIC EXAMPLE — not a real study or product.",
}


def validate(spec: dict) -> list[str]:
    problems = []
    for key in ("width_in", "height_in", "title"):
        if not spec.get(key):
            problems.append(f"missing required field: {key}")
    if not spec.get("columns"):
        problems.append("poster has no columns")
    if not spec.get("approval_status"):
        problems.append(
            "no approval_status. State the regulatory status of anything discussed."
        )
    for name, key, floor in [
        ("body", "body_pt", MIN_BODY_PT),
        ("heading", "heading_pt", MIN_HEADING_PT),
        ("title", "title_pt", MIN_TITLE_PT),
    ]:
        val = spec.get(key)
        if val is not None and val < floor:
            problems.append(
                f"{key}={val} is below the {floor}pt minimum for {name} text. "
                f"A reader at two metres cannot read it. Cut content instead of "
                f"shrinking type — this cannot be fixed at the venue."
            )
    return problems


def _have(module: str) -> bool:
    """Inline capability check, kept small so a copied directory still works."""
    try:
        __import__(module)
        return True
    except ImportError:
        return False


def render_html(spec: dict) -> str:
    """Tier 3: a printable HTML poster at the correct physical aspect ratio.

    This one arguably beats the .pptx: it prints correctly from any browser,
    the type sizes are set in real inches so the two-metre legibility rule is
    preserved exactly, and there is no dependency at all.
    """
    from html import escape

    W, H = float(spec["width_in"]), float(spec["height_in"])
    # Whole numbers read better in CSS: "48in", not "48.0in".
    Ws = f"{W:g}"
    Hs = f"{H:g}"
    title_pt = spec.get("title_pt", max(MIN_TITLE_PT, int(W * 1.7)))
    heading_pt = spec.get("heading_pt", max(MIN_HEADING_PT, int(W * 0.85)))
    body_pt = spec.get("body_pt", max(MIN_BODY_PT, int(W * 0.55)))

    cols = []
    for col in spec.get("columns", []):
        secs = []
        for sec in col.get("sections", []):
            body = escape(str(sec.get("body", ""))).replace("\n", "<br>")
            img = ""
            if sec.get("image"):
                img = f'<img src="{escape(str(sec["image"]))}" alt="">'
            secs.append(f"<section><h2>{escape(sec['heading'])}</h2>"
                        f"<p>{body}</p>{img}</section>")
        cols.append("<div class=\"col\">" + "".join(secs) + "</div>")

    takeaway = (f'<div class="takeaway">{escape(spec["takeaway"])}</div>'
                if spec.get("takeaway") else "")

    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<title>{escape(spec.get('title', 'Poster'))}</title>
<style>
  @page {{ size: {Ws}in {Hs}in; margin: 0; }}
  html, body {{ margin: 0; padding: 0; background: #fff; color: #1a1a1a;
    font-family: Helvetica, Arial, sans-serif; }}
  .poster {{ width: {Ws}in; height: {Hs}in; box-sizing: border-box;
    padding: {W * 0.02}in; display: flex; flex-direction: column; }}
  h1 {{ font-size: {title_pt}pt; color: #00518c; margin: 0 0 .18in; line-height: 1.05; }}
  .meta {{ font-size: {int(body_pt * 0.95)}pt; margin-bottom: .06in; }}
  .affil {{ font-size: {int(body_pt * 0.8)}pt; color: #555; margin-bottom: .16in; }}
  .takeaway {{ font-size: {int(heading_pt * 0.9)}pt; font-weight: 700;
    color: #00518c; border-top: 3px solid #00518c; border-bottom: 3px solid #00518c;
    padding: .12in 0; margin: 0 0 .24in; }}
  .cols {{ display: flex; gap: {W * 0.015}in; flex: 1; }}
  .col {{ flex: 1; }}
  h2 {{ font-size: {heading_pt}pt; color: #00518c; margin: 0 0 .08in; }}
  section {{ margin-bottom: .28in; }}
  p {{ font-size: {body_pt}pt; line-height: 1.35; margin: 0; }}
  img {{ max-width: 100%; margin-top: .12in; }}
  footer {{ font-size: {int(body_pt * 0.75)}pt; border-top: 1px solid #ccc;
    padding-top: .1in; display: flex; justify-content: space-between; gap: 1in; }}
  .draft {{ color: #993300; font-weight: 700; }}
</style></head><body><div class="poster">
<h1>{escape(spec.get('title', ''))}</h1>
<div class="meta">{escape(spec.get('authors', ''))}
  {(' · ' + escape(spec['abstract_number'])) if spec.get('abstract_number') else ''}</div>
<div class="affil">{escape(spec.get('affiliations', ''))}</div>
{takeaway}
<div class="cols">{''.join(cols)}</div>
<footer>
  <span>{escape(spec.get('approval_status', ''))}<br>
        <span style="color:#555">{escape(spec.get('footer', ''))}</span></span>
  <span class="draft">DRAFT — REQUIRES QUALIFIED MEDICAL REVIEW</span>
</footer>
</div></body></html>
"""


def build(spec: dict, out: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    W, H = float(spec["width_in"]), float(spec["height_in"])
    title_pt = spec.get("title_pt", max(MIN_TITLE_PT, int(W * 1.7)))
    heading_pt = spec.get("heading_pt", max(MIN_HEADING_PT, int(W * 0.85)))
    body_pt = spec.get("body_pt", max(MIN_BODY_PT, int(W * 0.55)))

    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    s = prs.slides.add_slide(prs.slide_layouts[6])

    INK = RGBColor(0x1A, 0x1A, 0x1A)
    ACCENT = RGBColor(0x00, 0x51, 0x8C)
    MUTED = RGBColor(0x55, 0x55, 0x55)
    WARN = RGBColor(0x99, 0x33, 0x00)

    def box(l, t, w, h, text, size, *, bold=False, colour=INK,
            align=PP_ALIGN.LEFT, space=6):
        tb = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        first = True
        for line in str(text).split("\n"):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = align
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = colour
            p.space_after = Pt(space)
        return tb

    margin = W * 0.02
    y = margin

    box(margin, y, W - 2 * margin, H * 0.10, spec["title"],
        title_pt, bold=True, colour=ACCENT)
    y += H * 0.095

    meta = " · ".join(x for x in [spec.get("authors", ""),
                                  spec.get("abstract_number", "")] if x)
    box(margin, y, W - 2 * margin, H * 0.035, meta, int(body_pt * 0.95))
    y += H * 0.035
    if spec.get("affiliations"):
        box(margin, y, W - 2 * margin, H * 0.03, spec["affiliations"],
            int(body_pt * 0.8), colour=MUTED)
        y += H * 0.032

    # A single takeaway at the top is the most effective layout decision
    # available: it is what a passer-by reads in three seconds.
    if spec.get("takeaway"):
        box(margin, y, W - 2 * margin, H * 0.07,
            spec["takeaway"], int(heading_pt * 0.9), bold=True, colour=ACCENT)
        y += H * 0.08

    top_of_columns = y
    n_cols = len(spec["columns"])
    gutter = W * 0.015
    col_w = (W - 2 * margin - gutter * (n_cols - 1)) / n_cols

    for i, col in enumerate(spec["columns"]):
        x = margin + i * (col_w + gutter)
        cy = top_of_columns
        for sec in col.get("sections", []):
            box(x, cy, col_w, H * 0.045, sec["heading"], heading_pt,
                bold=True, colour=ACCENT)
            cy += H * 0.05
            body = str(sec.get("body", ""))
            # Rough vertical allowance; deliberately generous so text is not
            # clipped, at the cost of leaving whitespace — which posters need.
            est_lines = sum(
                max(1, len(line) // max(20, int(col_w * 90 / body_pt)) + 1)
                for line in body.split("\n")
            )
            height = max(H * 0.05, est_lines * body_pt / 72 * 1.5)
            box(x, cy, col_w, height, body, body_pt)
            cy += height + H * 0.025
            if sec.get("image"):
                p = Path(sec["image"])
                if p.exists():
                    s.shapes.add_picture(str(p), Inches(x), Inches(cy),
                                         width=Inches(col_w))
                    cy += col_w * 0.62 + H * 0.02

    foot_y = H - margin - H * 0.035
    box(margin, foot_y, W - 2 * margin, H * 0.02,
        spec.get("approval_status", ""), int(body_pt * 0.75), colour=INK)
    box(margin, foot_y + H * 0.018, W - 2 * margin, H * 0.02,
        spec.get("footer", ""), int(body_pt * 0.7), colour=MUTED)
    box(W / 2, foot_y + H * 0.018, W / 2 - margin, H * 0.02,
        "DRAFT — REQUIRES QUALIFIED MEDICAL REVIEW",
        int(body_pt * 0.7), colour=WARN, align=PP_ALIGN.RIGHT)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return title_pt, heading_pt, body_pt


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--spec")
    ap.add_argument("--out", default="poster.pptx")
    ap.add_argument("--example", action="store_true")
    args = ap.parse_args()

    if args.example:
        print(json.dumps(EXAMPLE, indent=2))
        return 0
    if not args.spec:
        ap.print_help()
        return 2

    path = Path(args.spec)
    if not path.exists():
        print(f"No such file: {path}", file=sys.stderr)
        return 2
    spec = json.loads(path.read_text(encoding="utf-8"))

    problems = validate(spec)
    for p in problems:
        print(f"  BLOCKED  {p}")
    if problems:
        return 1

    PRINT_NOTE = (
        "\nBefore printing: confirm the dimensions with BOTH the congress and the\n"
        "printer, and check legibility by printing a page at 25% and reading it\n"
        "at arm's length. That approximates two metres from the full-size poster."
    )
    out = Path(args.out)

    if _have("pptx"):
        sizes = build(spec, out)
        print(f"Wrote {out} at {spec['width_in']}×{spec['height_in']} inches.")
        print(f"Type sizes: title {sizes[0]}pt · headings {sizes[1]}pt · body {sizes[2]}pt")
        print(PRINT_NOTE)
        return 0

    html_path = out.with_suffix(".html")
    html_path.parent.mkdir(parents=True, exist_ok=True)
    html_path.write_text(render_html(spec), encoding="utf-8")
    print(f"""⚠ DEGRADED OUTPUT — python-pptx is not available in this environment.

   Delivered:      {html_path}  ({spec['width_in']}×{spec['height_in']}in, print-ready)
   Not delivered:  {out}
   To get it:      pip install python-pptx
                   python3 {Path(__file__).name} --spec {args.spec} --out {out}

   The content is complete, and for a poster this format arguably beats the
   .pptx — it prints correctly from any browser at the exact physical size, and
   the minimum type sizes are preserved in real inches, so the two-metre
   legibility rule still holds. Print to PDF from the browser for the printer.
""")
    print(PRINT_NOTE)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
