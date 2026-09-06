#!/usr/bin/env python3
"""Read PDFs, Word files, decks, workbooks and CSVs — and say what failed.

Almost every real Medical Affairs task starts with a file somebody sends you,
and nothing else in this library can open one. This is the front door.

    S=skills/document-ingestion/scripts/ingest.py

    python3 $S --file paper.pdf
    python3 $S --file paper.pdf --tables
    python3 $S --file deck.pptx --format json
    python3 $S --file insights.xlsx --sheet 2
    python3 $S --dir ./uploads
    python3 $S --file transcript.docx --scan-safety

The governing rule: NEVER guess at content that could not be read. A plausible
summary of a document you could not open is the most damaging failure available
here, because the output looks exactly like a real one. This script reports the
extraction path it used and what it could not reach, so you can tell the
difference.

Every reader is optional. With none installed the script still runs, still reads
CSV and plain text, and tells you exactly what to install for the rest.
"""

from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
import zipfile
from dataclasses import dataclass, field, asdict
from pathlib import Path

# Patterns from medical-affairs-foundations. A detection aid, never a control.
AE_PATTERNS = [
    (r"\b(hospitali[sz]|admitted to hospital|inpatient)\w*", "possible serious AE (hospitalisation)"),
    (r"\b(died|death|fatal|mortality)\b", "possible fatal outcome"),
    (r"\b(pregnan|breastfeed|lactat)\w*", "special situation: pregnancy/lactation exposure"),
    (r"\b(overdose|double dose|took two|twice the dose)\b", "special situation: overdose"),
    (r"\b(medication error|wrong dose|dosing error|mix[- ]up)\b", "special situation: medication error"),
    (r"\b(off[- ]label|outside the label|unapproved (use|indication))\b", "special situation: off-label use"),
    (r"\b(grade [3-5]|serious adverse|SAE\b)", "possible serious AE"),
    (r"\b(discontinu\w+|stopped) (because|due to|owing to)\b", "AE leading to discontinuation"),
    (r"\b(lack of (efficacy|effect)|stopped working|no response)\b", "possible lack of therapeutic effect"),
    (r"\b(jamm\w+|leak\w+|malfunction\w*|would not (deliver|fire|prime)|defect\w*)\b",
     "possible product quality complaint"),
    (r"\b(adverse (event|reaction)|side effect|toxicit)\w*", "adverse event language"),
]

# Direct identifiers that must never be reproduced downstream.
PHI_HINTS = [
    (r"\b\d{3}-\d{2}-\d{4}\b", "US social security number format"),
    (r"\b(MRN|medical record (number|no))\b[:\s#]*\w+", "medical record number"),
    (r"\b\d{1,2}[/-]\d{1,2}[/-]\d{4}\b.{0,30}\b(DOB|date of birth)\b", "date of birth"),
    (r"\b(DOB|date of birth)\b[:\s]*\d", "date of birth"),
]


@dataclass
class Doc:
    path: str
    fmt: str = ""
    method: str = ""
    text: str = ""
    tables: list = field(default_factory=list)
    units: list = field(default_factory=list)   # pages / slides / sheets
    unread: list = field(default_factory=list)  # what could NOT be read
    notes: list = field(default_factory=list)


def _have(module: str) -> bool:
    try:
        __import__(module)
        return True
    except ImportError:
        return False


# --------------------------------------------------------------------------
# Readers
# --------------------------------------------------------------------------


def read_pdf(path: Path, want_tables: bool) -> Doc:
    d = Doc(path=str(path), fmt="pdf")

    if _have("pdfplumber"):
        import pdfplumber

        d.method = "pdfplumber (text layer + layout)"
        with pdfplumber.open(str(path)) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                t = page.extract_text() or ""
                if t.strip():
                    d.units.append({"page": i, "chars": len(t)})
                    d.text += f"\n\n--- page {i} ---\n{t}"
                else:
                    d.unread.append(
                        f"page {i}: no text layer — probably a scanned image. "
                        f"Its contents are NOT known."
                    )
                if want_tables:
                    for tbl in page.extract_tables() or []:
                        d.tables.append({"page": i, "rows": tbl})
        if d.tables:
            d.notes.append(
                "Table structure in a PDF is INFERRED, not stored. A misaligned "
                "column turns a hazard ratio into a CI bound. Verify any number "
                "that matters against the rendered document."
            )
        return d

    if _have("pypdf"):
        from pypdf import PdfReader

        d.method = "pypdf (embedded text layer only)"
        reader = PdfReader(str(path))
        for i, page in enumerate(reader.pages, 1):
            t = page.extract_text() or ""
            if t.strip():
                d.units.append({"page": i, "chars": len(t)})
                d.text += f"\n\n--- page {i} ---\n{t}"
            else:
                d.unread.append(
                    f"page {i}: no text layer — probably a scanned image. "
                    f"Its contents are NOT known."
                )
        if want_tables:
            d.notes.append(
                "Table extraction needs pdfplumber (pip install pdfplumber); "
                "pypdf returns flowed text only."
            )
        return d

    d.method = "none available"
    d.unread.append("the entire file — no PDF reader is installed")
    d.notes.append(
        "No PDF reader available. Install one:  pip install pypdf\n"
        "  (or pdfplumber, which also extracts tables)\n"
        "Alternatives that need no install: paste the relevant text; or give the\n"
        "DOI/PMID and retrieve the abstract via skills/pubmed-search."
    )
    return d


def read_docx(path: Path) -> Doc:
    d = Doc(path=str(path), fmt="docx")
    if not _have("docx"):
        d.method = "none available"
        d.unread.append("the entire file — python-docx is not installed")
        d.notes.append("pip install python-docx")
        return d
    from docx import Document

    d.method = "python-docx"
    doc = Document(str(path))
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            style = (p.style.name or "").lower()
            prefix = "#" * min(int(style[-1]), 6) + " " if style.startswith("heading") and style[-1].isdigit() else ""
            parts.append(prefix + p.text)
    d.text = "\n\n".join(parts)
    for i, tbl in enumerate(doc.tables, 1):
        rows = [[c.text.strip() for c in r.cells] for r in tbl.rows]
        d.tables.append({"table": i, "rows": rows})
    d.units = [{"paragraphs": len(doc.paragraphs), "tables": len(doc.tables)}]
    d.notes.append(
        "Tracked changes: this reads the document as stored. Confirm whether you "
        "have the accepted or the marked-up version before quoting it."
    )
    return d


def read_pptx(path: Path) -> Doc:
    d = Doc(path=str(path), fmt="pptx")
    if not _have("pptx"):
        d.method = "none available"
        d.unread.append("the entire file — python-pptx is not installed")
        d.notes.append("pip install python-pptx")
        return d
    from pptx import Presentation

    d.method = "python-pptx"
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides, 1):
        bits = []
        images = 0
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                bits.append(shape.text_frame.text.strip())
            if shape.shape_type is not None and "PICTURE" in str(shape.shape_type):
                images += 1
            if getattr(shape, "has_table", False):
                rows = [[c.text.strip() for c in r.cells] for r in shape.table.rows]
                d.tables.append({"slide": i, "rows": rows})
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        d.units.append({"slide": i, "shapes_with_text": len(bits), "images": images})
        d.text += f"\n\n--- slide {i} ---\n" + "\n".join(bits)
        if notes:
            d.text += f"\n[speaker notes] {notes}"
        if images:
            d.unread.append(
                f"slide {i}: {images} image(s) — any text inside them is NOT read"
            )
    d.notes.append(
        "Shape order does not always match reading order, especially with grouped "
        "shapes. Check the sequence if it matters."
    )
    return d


def read_xlsx(path: Path, sheet: str | None) -> Doc:
    d = Doc(path=str(path), fmt="xlsx")
    if not _have("openpyxl"):
        d.method = "none available"
        d.unread.append("the entire file — openpyxl is not installed")
        d.notes.append(
            "pip install openpyxl\n"
            "Alternative with no install: re-save the sheet as CSV and pass that."
        )
        return d
    import openpyxl

    d.method = "openpyxl (cached values)"
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    names = wb.sheetnames
    targets = names
    if sheet:
        targets = [names[int(sheet) - 1]] if sheet.isdigit() else [sheet]
    for name in targets:
        ws = wb[name]
        rows = [[("" if c is None else str(c)) for c in r] for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(x.strip() for x in r)]
        d.tables.append({"sheet": name, "rows": rows})
        d.units.append({"sheet": name, "rows": len(rows)})
    if len(names) > len(targets):
        d.notes.append(f"Sheets not read: {', '.join(n for n in names if n not in targets)}")
    d.notes.append(
        "Formulas are read as their CACHED value — if the workbook was never "
        "recalculated, those values may be stale. Merged cells and multi-row "
        "headers are not resolved."
    )
    return d


def read_csv(path: Path) -> Doc:
    d = Doc(path=str(path), fmt="csv")
    d.method = "csv (stdlib)"
    raw = path.read_text(encoding="utf-8", errors="replace")
    lines = raw.splitlines()
    # Leading comment blocks are common in exports and break header detection.
    skipped = 0
    while lines and lines[0].lstrip().startswith("#"):
        skipped += 1
        lines.pop(0)
    if skipped:
        d.notes.append(f"Skipped {skipped} leading comment line(s) before the header.")
    try:
        dialect = csv.Sniffer().sniff("\n".join(lines[:20]))
        delim = dialect.delimiter
    except csv.Error:
        delim = ","
        d.notes.append("Could not sniff the delimiter; assumed comma.")
    rows = list(csv.reader(lines, delimiter=delim))
    d.tables.append({"sheet": path.name, "rows": rows})
    d.units.append({"rows": len(rows), "delimiter": delim})
    d.text = raw
    return d


def read_text(path: Path) -> Doc:
    d = Doc(path=str(path), fmt="text", method="plain read")
    d.text = path.read_text(encoding="utf-8", errors="replace")
    d.units.append({"chars": len(d.text)})
    return d


READERS = {
    ".pdf": lambda p, a: read_pdf(p, a.tables),
    ".docx": lambda p, a: read_docx(p),
    ".pptx": lambda p, a: read_pptx(p),
    ".xlsx": lambda p, a: read_xlsx(p, a.sheet),
    ".xlsm": lambda p, a: read_xlsx(p, a.sheet),
    ".csv": lambda p, a: read_csv(p),
    ".tsv": lambda p, a: read_csv(p),
    ".txt": lambda p, a: read_text(p),
    ".md": lambda p, a: read_text(p),
    ".json": lambda p, a: read_text(p),
}


def safety_scan(text: str) -> list[dict]:
    hits = []
    for line in text.splitlines():
        s = line.strip()
        if not s:
            continue
        for pattern, label in AE_PATTERNS:
            if re.search(pattern, s, re.I):
                hits.append({"category": label, "verbatim": s[:400]})
                break
    return hits


def phi_scan(text: str) -> list[str]:
    found = []
    for pattern, label in PHI_HINTS:
        if re.search(pattern, text, re.I):
            found.append(label)
    return sorted(set(found))


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--file")
    ap.add_argument("--dir")
    ap.add_argument("--sheet", help="sheet name or 1-based index (xlsx)")
    ap.add_argument("--tables", action="store_true", help="attempt table extraction")
    ap.add_argument("--scan-safety", action="store_true",
                    help="run the AE/PQC/special-situation scan over the text")
    ap.add_argument("--format", default="text", choices=["text", "json"])
    ap.add_argument("--max-chars", type=int, default=20000,
                    help="truncate printed text (the full text is in --format json)")
    args = ap.parse_args()

    paths: list[Path] = []
    if args.file:
        paths = [Path(args.file)]
    elif args.dir:
        paths = sorted(p for p in Path(args.dir).rglob("*")
                       if p.is_file() and p.suffix.lower() in READERS)
    else:
        ap.print_help()
        return 2

    docs = []
    for p in paths:
        if not p.exists():
            print(f"No such file: {p}", file=sys.stderr)
            return 2
        reader = READERS.get(p.suffix.lower())
        if not reader:
            print(f"Unsupported format: {p.suffix} ({p.name})", file=sys.stderr)
            continue
        docs.append(reader(p, args))

    if args.format == "json":
        print(json.dumps([asdict(d) for d in docs], indent=2))
        return 0

    unreadable = 0
    for d in docs:
        print(f"\n{'='*72}\n{d.path}  [{d.fmt}]\n{'='*72}")
        print(f"Extraction path: {d.method}")
        if d.units:
            print(f"Read: {d.units}")

        if d.unread:
            unreadable += 1
            print("\n⚠ NOT READ:")
            for u in d.unread:
                print(f"   - {u}")
            print("   Do not infer what these contain. Say so in your output.")

        for n in d.notes:
            print(f"\nNote: {n}")

        phi = phi_scan(d.text)
        if phi:
            print("\n⚠ POSSIBLE PERSONAL/PATIENT IDENTIFIERS DETECTED")
            for f in phi:
                print(f"   - {f}")
            print("   Do NOT reproduce these downstream. Redact direct identifiers,")
            print("   and tell the user what you found before going further.")

        if args.scan_safety:
            hits = safety_scan(d.text)
            print(f"\n⚠ SAFETY SCAN — {len(hits)} potential finding(s)")
            for h in hits[:40]:
                print(f'   [{h["category"]}]  "{h["verbatim"]}"')
            if len(hits) > 40:
                print(f"   ... and {len(hits) - 40} more (use --format json for all)")
            print(
                "\n   Route every one to Pharmacovigilance through your own system,\n"
                "   with the verbatim intact. This is a keyword aid and WILL miss\n"
                "   cases; the human's reporting clock started when the document\n"
                "   reached them, not when this ran."
            )

        if d.tables:
            print(f"\nTables extracted: {len(d.tables)}")
            for t in d.tables[:3]:
                loc = t.get("page") or t.get("slide") or t.get("sheet") or t.get("table")
                print(f"  [{loc}] {len(t['rows'])} rows × "
                      f"{max((len(r) for r in t['rows']), default=0)} cols")
                for r in t["rows"][:3]:
                    print("     " + " | ".join(str(c)[:24] for c in r))

        if d.text.strip():
            body = d.text.strip()
            print(f"\n--- text ({len(body)} chars) ---")
            print(body[: args.max_chars])
            if len(body) > args.max_chars:
                print(f"\n[truncated at {args.max_chars} chars — "
                      f"use --format json for the full text]")

    if unreadable:
        print(
            f"\n\n{unreadable} document(s) had content that could not be read.\n"
            f"Never fill that gap with something plausible. Say what failed and\n"
            f"offer a route: paste the text, give a DOI/PMID for retrieval via\n"
            f"pubmed-search, or install the missing reader."
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
